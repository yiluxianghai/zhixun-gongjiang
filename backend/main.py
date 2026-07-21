"""
工程监理巡视问题分类闭环管理智能体
FastAPI 主应用

启动: uvicorn main:app --reload --host 0.0.0.0 --port 8000
"""
import json
import os
import uuid
from datetime import datetime, date, timedelta
from typing import Optional, List

from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy.orm import Session
import openpyxl

from database import (
    init_db, get_db, Project, InspectionProblem, RectificationRecord,
    OutputDocument, generate_problem_no, get_next_seq,
    AIModelConfig, KnowledgeBaseConfig, AnalysisSkill,
    KnowledgeDocument, DocumentChunk,
)
from ai_engine import analyze_problem, analyze_image
from rag_engine import process_document, retrieve_context, build_rag_context
from document_generator import (
    generate_notice, generate_inspection_record,
    generate_ledger, generate_analysis_report
)

app = FastAPI(
    title="工程监理巡视问题分类闭环管理智能体",
    description="围绕现场巡视问题的发现—识别—分类—整改—复查—销项全过程",
    version="1.0.0"
)

# CORS配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 静态文件服务 - 上传的图片
UPLOAD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)
app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")

# 知识库文档存储目录
KB_DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "kb_docs")
os.makedirs(KB_DOCS_DIR, exist_ok=True)


# ========== Pydantic 模型 ==========

class ProjectCreate(BaseModel):
    project_code: str
    project_name: str


class ProblemAnalyzeRequest(BaseModel):
    project_id: int
    inspection_area: str
    inspector: str
    inspection_date: str
    raw_description: str
    model_id: Optional[int] = None
    kb_id: Optional[int] = None
    skill_id: Optional[int] = None

class AIModelCreate(BaseModel):
    name: str
    provider: str = "openai"
    api_key: str = ""
    base_url: str = "https://api.openai.com/v1"
    model_name: str = "gpt-4o-mini"
    temperature: float = 0.3
    max_tokens: int = 2000

class KBCreate(BaseModel):
    name: str
    description: str = ""
    content: str = "{}"

class SkillCreate(BaseModel):
    name: str
    description: str = ""
    system_prompt: str = ""
    user_prompt_template: str = ""


class ProblemCreateRequest(BaseModel):
    project_id: int
    inspection_area: str
    inspector: str
    inspection_date: str
    raw_description: str
    # AI分析结果（可由前端修改后提交）
    standardized_desc: str
    category_code: str
    category_name: str
    specialty_code: str
    specialty_name: str
    risk_level: str
    risk_reason: Optional[str] = ""
    rectification_req: str
    rectification_deadline_days: int = 7
    review_points: str
    responsible_party: str
    confidence: float = 0.0
    photo_urls: Optional[str] = ""


class StatusUpdateRequest(BaseModel):
    status: str
    operator: Optional[str] = ""


class RectificationFeedbackRequest(BaseModel):
    feedback: str
    operator: str


class ReviewRequest(BaseModel):
    result: str  # 通过/不通过
    review_comment: Optional[str] = ""
    operator: str


# ========== 启动事件 ==========

@app.on_event("startup")
def startup():
    init_db()
    print("[系统] 工程监理巡视问题分类闭环管理智能体已启动")
    print("[系统] API文档: http://localhost:8000/docs")


# ========== 项目管理 ==========

@app.post("/api/projects")
def create_project(req: ProjectCreate, db: Session = Depends(get_db)):
    """创建项目"""
    project = Project(project_code=req.project_code, project_name=req.project_name)
    db.add(project)
    db.commit()
    db.refresh(project)
    return {"code": 200, "data": {"id": project.id, "project_code": project.project_code, "project_name": project.project_name}}


@app.get("/api/projects")
def list_projects(db: Session = Depends(get_db)):
    """获取项目列表"""
    projects = db.query(Project).all()
    result = []
    for p in projects:
        problem_count = db.query(InspectionProblem).filter(InspectionProblem.project_id == p.id).count()
        result.append({
            "id": p.id,
            "project_code": p.project_code,
            "project_name": p.project_name,
            "problem_count": problem_count,
            "created_at": p.created_at.isoformat() if p.created_at else "",
        })
    return {"code": 200, "data": result}


# ========== 问题分析 ==========

@app.post("/api/problems/analyze")
async def analyze_problem_api(req: ProblemAnalyzeRequest, db: Session = Depends(get_db)):
    """AI智能分析问题（核心接口，支持动态配置 + RAG检索）"""
    # 获取模型配置
    model_config = None
    if req.model_id:
        m = db.query(AIModelConfig).filter(AIModelConfig.id == req.model_id).first()
        if m:
            model_config = {"provider": m.provider, "api_key": m.api_key,
                            "base_url": m.base_url, "model_name": m.model_name,
                            "temperature": m.temperature}
    else:
        m = db.query(AIModelConfig).filter(AIModelConfig.is_active == 1).first()
        if m:
            model_config = {"provider": m.provider, "api_key": m.api_key,
                            "base_url": m.base_url, "model_name": m.model_name,
                            "temperature": m.temperature}

    # 获取知识库配置
    kb_content = None
    kb_id = req.kb_id
    if not kb_id:
        active_kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.is_active == 1).first()
        kb_id = active_kb.id if active_kb else None
    if kb_id:
        kb_obj = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
        if kb_obj:
            try:
                kb_content = json.loads(kb_obj.content)
            except:
                kb_content = None

    # 获取分析技能
    skill = None
    skill_id = req.skill_id
    if not skill_id:
        active_skill = db.query(AnalysisSkill).filter(AnalysisSkill.is_active == 1).first()
        skill_id = active_skill.id if active_skill else None
    if skill_id:
        s = db.query(AnalysisSkill).filter(AnalysisSkill.id == skill_id).first()
        if s:
            skill = {"system_prompt": s.system_prompt, "user_prompt_template": s.user_prompt_template}

    # RAG检索：从知识库文档中检索相关内容
    rag_context = None
    if kb_id:
        # 查询该知识库下所有已处理完成的文档分块
        docs = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.kb_id == kb_id,
            KnowledgeDocument.status == "ready",
        ).all()
        
        if docs:
            doc_ids = [d.id for d in docs]
            chunks = db.query(DocumentChunk).filter(DocumentChunk.doc_id.in_(doc_ids)).all()
            
            if chunks:
                # 构建查询文本
                query_text = f"{req.raw_description} {req.inspection_area}"
                
                # 获取用于嵌入的API配置
                emb_api_key = model_config.get("api_key", "") if model_config else ""
                emb_base_url = model_config.get("base_url", "") if model_config else ""
                
                chunks_data = [{"content": c.content, "embedding": c.embedding} for c in chunks]
                retrieved = await retrieve_context(
                    query_text, chunks_data,
                    api_key=emb_api_key, base_url=emb_base_url,
                    top_k=5,
                )
                rag_context = build_rag_context(retrieved)
                print(f"[RAG] 检索完成，上下文长度: {len(rag_context) if rag_context else 0} 字符")

    result = await analyze_problem(
        req.raw_description, req.inspection_area,
        model_config=model_config, kb_content=kb_content, skill=skill,
        rag_context=rag_context,
    )
    return {"code": 200, "data": result}


# ========== 问题管理 ==========

@app.post("/api/problems")
def create_problem(req: ProblemCreateRequest, db: Session = Depends(get_db)):
    """创建问题（确认后保存）"""
    project = db.query(Project).filter(Project.id == req.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # 生成编号
    seq = get_next_seq(db, req.project_id)
    problem_no = generate_problem_no(project.project_code, req.specialty_code, seq)

    # 计算整改期限
    deadline = (datetime.now() + timedelta(days=req.rectification_deadline_days)).date()

    problem = InspectionProblem(
        problem_no=problem_no,
        project_id=req.project_id,
        inspection_date=datetime.strptime(req.inspection_date, "%Y-%m-%d").date(),
        inspection_area=req.inspection_area,
        inspector=req.inspector,
        raw_description=req.raw_description,
        standardized_desc=req.standardized_desc,
        category_code=req.category_code,
        category_name=req.category_name,
        specialty_code=req.specialty_code,
        specialty_name=req.specialty_name,
        risk_level=req.risk_level,
        risk_reason=req.risk_reason,
        rectification_req=req.rectification_req,
        rectification_deadline=deadline,
        review_points=req.review_points,
        responsible_party=req.responsible_party,
        ai_confidence=req.confidence,
        photo_urls=req.photo_urls,
        status="待整改",
    )
    db.add(problem)
    db.commit()
    db.refresh(problem)

    # 添加整改通知记录
    record = RectificationRecord(
        problem_id=problem.id,
        record_type="整改通知",
        content=f"问题已创建并下发整改通知，整改期限：{deadline.isoformat()}",
        operator=req.inspector,
    )
    db.add(record)
    db.commit()

    return {"code": 200, "data": {"id": problem.id, "problem_no": problem.problem_no}}


# ========== 图片上传 ==========

@app.post("/api/ai/analyze-image")
async def analyze_image_api(
    file: UploadFile = File(...),
    inspection_area: str = Form(default=""),
    db: Session = Depends(get_db),
):
    """AI图片识别：分析工程现场照片，检测安全违规和质量缺陷"""
    # 验证文件类型
    allowed_types = ["image/jpeg", "image/png", "image/webp"]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail=f"不支持的图片类型: {file.content_type}，仅支持: {', '.join(allowed_types)}")

    # 读取图片并转base64
    image_content = await file.read()
    import base64
    image_base64 = base64.b64encode(image_content).decode('utf-8')

    # 获取激活的模型配置
    m = db.query(AIModelConfig).filter(AIModelConfig.is_active == 1).first()
    if not m:
        return {"code": 200, "data": {
            "has_issues": False,
            "violations": [],
            "defects": [],
            "risk_level": "一般",
            "description": "未配置AI模型，无法进行图片识别。请在AI配置中添加支持视觉的模型（如GPT-4o）。",
            "recommendations": ["前往AI配置页面添加并激活支持视觉的LLM模型"],
            "confidence": 0.0,
        }}

    model_config = {
        "provider": m.provider,
        "api_key": m.api_key,
        "base_url": m.base_url,
        "model_name": m.model_name,
        "temperature": m.temperature,
    }

    result = await analyze_image(
        image_base64=image_base64,
        image_mime=file.content_type,
        inspection_area=inspection_area,
        model_config=model_config,
    )

    # 如果LLM不支持视觉分析，尝试本地YOLO模型作为后备
    if result is None and m.provider != 'yolo_local':
        try:
            from yolo_detector import is_model_available, detect_image
            import base64 as b64mod
            if is_model_available():
                image_bytes = b64mod.b64decode(image_base64)
                yolo_result = detect_image(image_bytes, inspection_area)
                if yolo_result:
                    yolo_result['description'] = f"[本地YOLO模型识别]\n{yolo_result.get('description', '')}"
                    return {"code": 200, "data": yolo_result}
        except Exception as e:
            print(f"[图片识别] YOLO后备失败: {e}")

    if result:
        return {"code": 200, "data": result}
    else:
        return {"code": 200, "data": {
            "has_issues": False,
            "violations": [],
            "defects": [],
            "risk_level": "一般",
            "description": "图片识别失败。可能的原因：\n1. 当前模型不支持视觉分析（如DeepSeek）\n2. 未配置API Key\n\n解决方案：\n- 激活支持视觉的LLM模型（GPT-4o/智谱GLM-4V）\n- 或训练本地YOLO模型进行离线识别",
            "recommendations": ["在AI配置中激活支持视觉的LLM模型（GPT-4o/智谱GLM-4V）", "或训练本地YOLO模型：运行 backend/train/train_yolo.py"],
            "confidence": 0.0,
        }}


@app.post("/api/upload/photo")
async def upload_photo(file: UploadFile = File(...)):
    """上传现场照片"""
    # 验证文件类型
    allowed_types = ["image/jpeg", "image/png", "image/webp", "image/gif"]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {file.content_type}，仅支持: {', '.join(allowed_types)}")

    # 生成唯一文件名
    ext = file.filename.rsplit(".", 1)[-1] if "." in file.filename else "jpg"
    filename = f"{uuid.uuid4().hex}.{ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)

    # 保存文件
    content = await file.read()
    with open(filepath, "wb") as f:
        f.write(content)

    # 返回可访问的URL
    url = f"/uploads/{filename}"
    print(f"[上传] 图片已保存: {filename} ({len(content)} bytes)")

    return {"code": 200, "data": {"url": url, "filename": filename, "size": len(content)}}


# ========== Excel批量导入 ==========

@app.post("/api/problems/import-excel")
async def import_excel(
    project_id: int = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Excel批量导入问题（自动AI分析）"""
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # 读取Excel
    content = await file.read()
    try:
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Excel解析失败: {e}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="Excel无有效数据行")

    # 识别表头
    header = [str(c or "").strip() for c in rows[0]]
    # 支持的列名映射
    col_map = {}
    for i, h in enumerate(header):
        if "区域" in h or "部位" in h:
            col_map["area"] = i
        elif "巡视人" in h or "发现人" in h:
            col_map["inspector"] = i
        elif "日期" in h or "时间" in h:
            col_map["date"] = i
        elif "描述" in h or "问题" in h:
            col_map["desc"] = i

    # 如果未识别到表头，按固定列顺序处理
    if "desc" not in col_map:
        col_map = {"area": 0, "inspector": 1, "date": 2, "desc": 3}
        data_rows = rows  # 无表头，从第一行开始
    else:
        data_rows = rows[1:]  # 跳过表头

    results = []
    success_count = 0
    fail_count = 0

    # ===== 加载 AI 配置（模型 / 知识库 / 技能 / RAG文档分块）=====
    m = db.query(AIModelConfig).filter(AIModelConfig.is_active == 1).first()
    model_config = None
    if m:
        model_config = {"provider": m.provider, "api_key": m.api_key,
                        "base_url": m.base_url, "model_name": m.model_name,
                        "temperature": m.temperature}

    kb_content = None
    active_kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.is_active == 1).first()
    kb_id = active_kb.id if active_kb else None
    if kb_id:
        kb_obj = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
        if kb_obj:
            try:
                kb_content = json.loads(kb_obj.content)
            except:
                kb_content = None

    skill = None
    active_skill = db.query(AnalysisSkill).filter(AnalysisSkill.is_active == 1).first()
    if active_skill:
        skill = {"system_prompt": active_skill.system_prompt,
                 "user_prompt_template": active_skill.user_prompt_template}

    # 预加载 RAG 文档分块（所有行共享，检索在循环内按行查询）
    rag_chunks_data = []
    if kb_id:
        docs = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.kb_id == kb_id,
            KnowledgeDocument.status == "ready",
        ).all()
        if docs:
            doc_ids = [d.id for d in docs]
            chunks = db.query(DocumentChunk).filter(DocumentChunk.doc_id.in_(doc_ids)).all()
            rag_chunks_data = [{"content": c.content, "embedding": c.embedding} for c in chunks]
            print(f"[Excel导入] RAG文档已加载: {len(docs)}篇文档, {len(rag_chunks_data)}个分块")

    emb_api_key = model_config.get("api_key", "") if model_config else ""
    emb_base_url = model_config.get("base_url", "") if model_config else ""

    for idx, row in enumerate(data_rows):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        try:
            area = str(row[col_map["area"]] or "").strip() if col_map.get("area", -1) < len(row) else ""
            inspector_name = str(row[col_map["inspector"]] or "").strip() if col_map.get("inspector", -1) < len(row) else ""
            date_str = str(row[col_map["date"]] or "").strip() if col_map.get("date", -1) < len(row) else ""
            desc = str(row[col_map["desc"]] or "").strip() if col_map.get("desc", -1) < len(row) else ""

            if not desc:
                fail_count += 1
                results.append({"row": idx + 2, "status": "fail", "reason": "问题描述为空"})
                continue

            # 解析日期
            if date_str:
                try:
                    parsed_date = datetime.strptime(date_str[:10], "%Y-%m-%d").date()
                except ValueError:
                    try:
                        parsed_date = datetime.strptime(date_str[:10], "%Y/%m/%d").date()
                    except ValueError:
                        parsed_date = date.today()
            else:
                parsed_date = date.today()

            # RAG 检索：针对每行问题描述检索相关知识库片段
            rag_context = None
            if rag_chunks_data:
                query_text = f"{desc} {area}"
                retrieved = await retrieve_context(
                    query_text, rag_chunks_data,
                    api_key=emb_api_key, base_url=emb_base_url,
                    top_k=5,
                )
                rag_context = build_rag_context(retrieved)

            # AI分析（携带完整配置 + RAG上下文）
            analysis = await analyze_problem(
                desc, area,
                model_config=model_config, kb_content=kb_content,
                skill=skill, rag_context=rag_context,
            )

            # 生成编号
            seq = get_next_seq(db, project_id)
            problem_no = generate_problem_no(project.project_code, analysis.get("specialty_code", "CIV"), seq)
            deadline = (datetime.now() + timedelta(days=analysis.get("rectification_deadline_days", 7))).date()

            # 创建问题
            problem = InspectionProblem(
                problem_no=problem_no,
                project_id=project_id,
                inspection_date=parsed_date,
                inspection_area=area or "未指定",
                inspector=inspector_name or "批量导入",
                raw_description=desc,
                standardized_desc=analysis.get("standardized_description", desc),
                category_code=analysis.get("category_code", "QM"),
                category_name=analysis.get("category_name", "质量管理"),
                specialty_code=analysis.get("specialty_code", "CIV"),
                specialty_name=analysis.get("specialty_name", "土建"),
                risk_level=analysis.get("risk_level", "一般"),
                risk_reason=analysis.get("risk_reason", ""),
                rectification_req=analysis.get("rectification_req", ""),
                rectification_deadline=deadline,
                review_points=analysis.get("review_points", ""),
                responsible_party=analysis.get("responsible_party", "总承包单位"),
                ai_confidence=analysis.get("confidence", 0.65),
                status="待整改",
            )
            db.add(problem)
            db.commit()
            db.refresh(problem)

            # 添加整改通知记录
            record = RectificationRecord(
                problem_id=problem.id,
                record_type="整改通知",
                content=f"问题已创建（Excel批量导入），整改期限：{deadline.isoformat()}",
                operator=inspector_name or "批量导入",
            )
            db.add(record)
            db.commit()

            success_count += 1
            results.append({
                "row": idx + 2,
                "status": "success",
                "problem_no": problem.problem_no,
                "area": area,
                "desc": desc[:30],
                "category": analysis.get("category_name", ""),
                "risk_level": analysis.get("risk_level", ""),
            })
            print(f"[Excel导入] 第{idx + 2}行: {problem.problem_no} 创建成功")

        except Exception as e:
            fail_count += 1
            results.append({"row": idx + 2, "status": "fail", "reason": str(e)})
            print(f"[Excel导入] 第{idx + 2}行导入失败: {e}")

    return {
        "code": 200,
        "data": {
            "total": len(results),
            "success": success_count,
            "fail": fail_count,
            "results": results,
        }
    }


@app.get("/api/problems")
def list_problems(
    project_id: int = Query(...),
    status: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    risk_level: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """获取问题列表（支持筛选）"""
    query = db.query(InspectionProblem).filter(InspectionProblem.project_id == project_id)

    if status:
        query = query.filter(InspectionProblem.status == status)
    if category:
        query = query.filter(InspectionProblem.category_code == category)
    if risk_level:
        query = query.filter(InspectionProblem.risk_level == risk_level)
    if keyword:
        query = query.filter(InspectionProblem.raw_description.contains(keyword))

    problems = query.order_by(InspectionProblem.created_at.desc()).all()

    result = []
    for p in problems:
        # 检查是否超期
        is_overdue = False
        if p.status in ["待整改", "整改中"] and p.rectification_deadline:
            if p.rectification_deadline < date.today():
                is_overdue = True

        result.append({
            "id": p.id,
            "problem_no": p.problem_no,
            "inspection_date": p.inspection_date.isoformat() if p.inspection_date else "",
            "inspection_area": p.inspection_area,
            "inspector": p.inspector,
            "raw_description": p.raw_description,
            "standardized_desc": p.standardized_desc,
            "category_code": p.category_code,
            "category_name": p.category_name,
            "specialty_code": p.specialty_code,
            "specialty_name": p.specialty_name,
            "risk_level": p.risk_level,
            "risk_reason": p.risk_reason,
            "rectification_req": p.rectification_req,
            "rectification_deadline": p.rectification_deadline.isoformat() if p.rectification_deadline else "",
            "review_points": p.review_points,
            "responsible_party": p.responsible_party,
            "ai_confidence": p.ai_confidence,
            "status": p.status,
            "is_overdue": is_overdue,
            "rectification_feedback": p.rectification_feedback,
            "rectification_date": p.rectification_date.isoformat() if p.rectification_date else "",
            "review_date": p.review_date.isoformat() if p.review_date else "",
            "review_result": p.review_result,
            "created_at": p.created_at.isoformat() if p.created_at else "",
        })

    return {"code": 200, "data": result}


@app.get("/api/problems/{problem_id}")
def get_problem(problem_id: int, db: Session = Depends(get_db)):
    """获取问题详情"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    records = db.query(RectificationRecord).filter(
        RectificationRecord.problem_id == problem_id
    ).order_by(RectificationRecord.created_at).all()

    return {
        "code": 200,
        "data": {
            "id": problem.id,
            "problem_no": problem.problem_no,
            "inspection_date": problem.inspection_date.isoformat() if problem.inspection_date else "",
            "inspection_area": problem.inspection_area,
            "inspector": problem.inspector,
            "raw_description": problem.raw_description,
            "standardized_desc": problem.standardized_desc,
            "category_code": problem.category_code,
            "category_name": problem.category_name,
            "specialty_code": problem.specialty_code,
            "specialty_name": problem.specialty_name,
            "risk_level": problem.risk_level,
            "risk_reason": problem.risk_reason,
            "rectification_req": problem.rectification_req,
            "rectification_deadline": problem.rectification_deadline.isoformat() if problem.rectification_deadline else "",
            "review_points": problem.review_points,
            "responsible_party": problem.responsible_party,
            "ai_confidence": problem.ai_confidence,
            "status": problem.status,
            "rectification_feedback": problem.rectification_feedback,
            "rectification_date": problem.rectification_date.isoformat() if problem.rectification_date else "",
            "review_date": problem.review_date.isoformat() if problem.review_date else "",
            "review_result": problem.review_result,
            "created_at": problem.created_at.isoformat() if problem.created_at else "",
            "records": [
                {
                    "id": r.id,
                    "record_type": r.record_type,
                    "content": r.content,
                    "operator": r.operator,
                    "created_at": r.created_at.isoformat() if r.created_at else "",
                }
                for r in records
            ],
        }
    }


@app.put("/api/problems/{problem_id}/status")
def update_status(problem_id: int, req: StatusUpdateRequest, db: Session = Depends(get_db)):
    """更新问题状态"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    valid_statuses = ["待整改", "整改中", "待复查", "已销项", "未通过"]
    if req.status not in valid_statuses:
        raise HTTPException(status_code=400, detail=f"无效状态，可选: {valid_statuses}")

    old_status = problem.status
    problem.status = req.status
    problem.updated_at = datetime.utcnow()

    # 添加记录
    record = RectificationRecord(
        problem_id=problem_id,
        record_type="状态变更",
        content=f"状态从「{old_status}」变更为「{req.status}」",
        operator=req.operator or "系统",
    )
    db.add(record)
    db.commit()

    return {"code": 200, "data": {"id": problem.id, "status": problem.status}}


@app.post("/api/problems/{problem_id}/rectification")
def submit_rectification(problem_id: int, req: RectificationFeedbackRequest, db: Session = Depends(get_db)):
    """提交整改反馈"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    problem.rectification_feedback = req.feedback
    problem.rectification_date = date.today()
    problem.status = "待复查"
    problem.updated_at = datetime.utcnow()

    record = RectificationRecord(
        problem_id=problem_id,
        record_type="整改反馈",
        content=req.feedback,
        operator=req.operator,
    )
    db.add(record)
    db.commit()

    return {"code": 200, "data": {"id": problem.id, "status": problem.status}}


@app.post("/api/problems/{problem_id}/review")
def submit_review(problem_id: int, req: ReviewRequest, db: Session = Depends(get_db)):
    """提交复查结果"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    problem.review_date = date.today()
    problem.review_result = req.result

    if req.result == "通过":
        problem.status = "已销项"
    else:
        problem.status = "待整改"  # 退回整改

    problem.updated_at = datetime.utcnow()

    record = RectificationRecord(
        problem_id=problem_id,
        record_type="复查记录",
        content=f"复查结果：{req.result}。{req.review_comment}",
        operator=req.operator,
    )
    db.add(record)
    db.commit()

    return {"code": 200, "data": {"id": problem.id, "status": problem.status}}


@app.delete("/api/problems/{problem_id}")
def delete_problem(problem_id: int, db: Session = Depends(get_db)):
    """删除问题及其所有关联记录"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    problem_no = problem.problem_no
    # records 已通过 cascade="all, delete-orphan" 级联删除
    db.delete(problem)
    db.commit()

    print(f"[台账] 问题 {problem_no} (ID={problem_id}) 已删除")
    return {"code": 200, "data": {"id": problem_id, "problem_no": problem_no}}


# ========== 统计看板 ==========

@app.get("/api/statistics/dashboard")
def get_dashboard(project_id: int = Query(...), db: Session = Depends(get_db)):
    """获取统计看板数据"""
    problems = db.query(InspectionProblem).filter(
        InspectionProblem.project_id == project_id
    ).all()

    total = len(problems)
    status_count = {"待整改": 0, "整改中": 0, "待复查": 0, "已销项": 0, "未通过": 0}
    category_count = {}
    risk_count = {"一般": 0, "较大": 0, "重大": 0}
    overdue_count = 0

    for p in problems:
        status_count[p.status] = status_count.get(p.status, 0) + 1
        cat = p.category_name or "未分类"
        category_count[cat] = category_count.get(cat, 0) + 1
        if p.risk_level in risk_count:
            risk_count[p.risk_level] += 1
        if p.status in ["待整改", "整改中"] and p.rectification_deadline:
            if p.rectification_deadline < date.today():
                overdue_count += 1

    closure_rate = round(status_count["已销项"] / total * 100, 1) if total > 0 else 0

    return {
        "code": 200,
        "data": {
            "total": total,
            "status_count": status_count,
            "category_count": category_count,
            "risk_count": risk_count,
            "overdue_count": overdue_count,
            "closure_rate": closure_rate,
        }
    }


# ========== 文档生成 ==========

@app.post("/api/documents/notice/{problem_id}")
def gen_notice(problem_id: int, db: Session = Depends(get_db)):
    """生成监理通知单"""
    problem = db.query(InspectionProblem).filter(InspectionProblem.id == problem_id).first()
    if not problem:
        raise HTTPException(status_code=404, detail="问题不存在")

    project = db.query(Project).filter(Project.id == problem.project_id).first()
    project_name = project.project_name if project else "未知项目"

    result = generate_notice(problem, project_name)
    return {"code": 200, "data": result}


@app.post("/api/documents/inspection-record")
def gen_inspection_record(
    project_id: int = Query(...),
    inspection_date: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """生成巡视记录"""
    query = db.query(InspectionProblem).filter(InspectionProblem.project_id == project_id)
    if inspection_date:
        query = query.filter(InspectionProblem.inspection_date == datetime.strptime(inspection_date, "%Y-%m-%d").date())

    problems = query.order_by(InspectionProblem.inspection_date.desc()).all()
    if not problems:
        return {"code": 200, "data": {"type": "巡视记录", "title": "巡视记录", "content": {"message": "暂无巡视数据"}}}

    project = db.query(Project).filter(Project.id == project_id).first()
    project_name = project.project_name if project else "未知项目"

    result = generate_inspection_record(problems, project_name, inspection_date)
    return {"code": 200, "data": result}


@app.get("/api/documents/ledger")
def gen_ledger(project_id: int = Query(...), db: Session = Depends(get_db)):
    """生成问题闭环台账"""
    problems = db.query(InspectionProblem).filter(
        InspectionProblem.project_id == project_id
    ).order_by(InspectionProblem.created_at.desc()).all()

    result = generate_ledger(problems)
    return {"code": 200, "data": result}


@app.post("/api/documents/analysis-report")
def gen_analysis_report(project_id: int = Query(...), db: Session = Depends(get_db)):
    """生成问题分析报告"""
    problems = db.query(InspectionProblem).filter(
        InspectionProblem.project_id == project_id
    ).all()

    project = db.query(Project).filter(Project.id == project_id).first()
    project_name = project.project_name if project else "未知项目"

    result = generate_analysis_report(problems, project_name)
    return {"code": 200, "data": result}


# ========== AI配置管理 ==========

@app.get("/api/ai/config")
def get_ai_config(db: Session = Depends(get_db)):
    """获取当前启用的AI配置"""
    model = db.query(AIModelConfig).filter(AIModelConfig.is_active == 1).first()
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.is_active == 1).first()
    skill = db.query(AnalysisSkill).filter(AnalysisSkill.is_active == 1).first()
    return {"code": 200, "data": {
        "model": {"id": model.id, "name": model.name, "provider": model.provider, "model_name": model.model_name} if model else None,
        "knowledge_base": {"id": kb.id, "name": kb.name} if kb else None,
        "skill": {"id": skill.id, "name": skill.name} if skill else None,
    }}

# --- AI模型配置 ---

@app.get("/api/ai/models")
def list_models(db: Session = Depends(get_db)):
    models = db.query(AIModelConfig).order_by(AIModelConfig.created_at.desc()).all()
    return {"code": 200, "data": [{"id": m.id, "name": m.name, "provider": m.provider,
        "api_key": m.api_key[:8] + "****" if m.api_key and len(m.api_key) > 8 else "",
        "base_url": m.base_url, "model_name": m.model_name,
        "temperature": m.temperature, "max_tokens": m.max_tokens,
        "is_active": bool(m.is_active), "created_at": str(m.created_at)} for m in models]}

@app.post("/api/ai/models")
def create_model(req: AIModelCreate, db: Session = Depends(get_db)):
    model = AIModelConfig(**req.dict())
    db.add(model)
    db.commit()
    db.refresh(model)
    return {"code": 200, "data": {"id": model.id, "name": model.name}}

@app.put("/api/ai/models/{model_id}")
def update_model(model_id: int, req: AIModelCreate, db: Session = Depends(get_db)):
    model = db.query(AIModelConfig).filter(AIModelConfig.id == model_id).first()
    if not model:
        raise HTTPException(status_code=404, detail="模型配置不存在")
    for k, v in req.dict().items():
        setattr(model, k, v)
    db.commit()
    return {"code": 200, "data": {"id": model.id}}

@app.delete("/api/ai/models/{model_id}")
def delete_model(model_id: int, db: Session = Depends(get_db)):
    model = db.query(AIModelConfig).filter(AIModelConfig.id == model_id).first()
    if not model:
        raise HTTPException(status_code=404, detail="模型配置不存在")
    db.delete(model)
    db.commit()
    return {"code": 200, "data": {"id": model_id}}

@app.post("/api/ai/models/{model_id}/activate")
def activate_model(model_id: int, db: Session = Depends(get_db)):
    db.query(AIModelConfig).update({AIModelConfig.is_active: 0})
    model = db.query(AIModelConfig).filter(AIModelConfig.id == model_id).first()
    if not model:
        raise HTTPException(status_code=404, detail="模型配置不存在")
    model.is_active = 1
    db.commit()
    return {"code": 200, "data": {"id": model_id}}

# --- 本地YOLO模型 ---

@app.get("/api/ai/yolo-status")
def yolo_status():
    """检查本地YOLO模型是否可用"""
    try:
        from yolo_detector import is_model_available, get_model_info
        info = get_model_info()
        return {"code": 200, "data": info}
    except ImportError:
        return {"code": 200, "data": {
            "available": False,
            "error": "yolo_detector模块未安装",
        }}

@app.post("/api/ai/yolo/activate")
def activate_yolo(db: Session = Depends(get_db)):
    """注册并激活本地YOLO模型"""
    try:
        from yolo_detector import is_model_available, get_model_info
    except ImportError:
        raise HTTPException(status_code=400, detail="YOLO模块未安装，请确认conda环境和ultralytics已配置")

    if not is_model_available():
        raise HTTPException(status_code=400, detail="未找到YOLO模型文件，请确认models/engine_safety.pt已就位")

    info = get_model_info()

    # 查找或创建YOLO模型配置
    yolo_model = db.query(AIModelConfig).filter(AIModelConfig.provider == "yolo_local").first()
    if yolo_model:
        yolo_model.name = "本地YOLO模型"
        yolo_model.model_name = "engine_safety.pt"
    else:
        yolo_model = AIModelConfig(
            name="本地YOLO模型",
            provider="yolo_local",
            api_key="",
            base_url="",
            model_name="engine_safety.pt",
            temperature=0.0,
            max_tokens=0,
        )
        db.add(yolo_model)

    # 激活YOLO模型
    db.query(AIModelConfig).update({AIModelConfig.is_active: 0})
    yolo_model.is_active = 1
    db.commit()
    db.refresh(yolo_model)

    return {"code": 200, "data": {
        "id": yolo_model.id,
        "name": yolo_model.name,
        "model_info": info,
    }}

# --- 知识库配置 ---

@app.get("/api/ai/knowledge-bases")
def list_kbs(db: Session = Depends(get_db)):
    kbs = db.query(KnowledgeBaseConfig).order_by(KnowledgeBaseConfig.created_at.desc()).all()
    return {"code": 200, "data": [{"id": k.id, "name": k.name, "description": k.description,
        "is_active": bool(k.is_active), "created_at": str(k.created_at)} for k in kbs]}

@app.post("/api/ai/knowledge-bases")
def create_kb(req: KBCreate, db: Session = Depends(get_db)):
    kb = KnowledgeBaseConfig(**req.dict())
    db.add(kb)
    db.commit()
    db.refresh(kb)
    return {"code": 200, "data": {"id": kb.id, "name": kb.name}}

@app.get("/api/ai/knowledge-bases/{kb_id}")
def get_kb(kb_id: int, db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    return {"code": 200, "data": {"id": kb.id, "name": kb.name, "description": kb.description,
        "content": kb.content, "is_active": bool(kb.is_active)}}

@app.put("/api/ai/knowledge-bases/{kb_id}")
def update_kb(kb_id: int, req: KBCreate, db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    for k, v in req.dict().items():
        setattr(kb, k, v)
    db.commit()
    return {"code": 200, "data": {"id": kb.id}}

@app.delete("/api/ai/knowledge-bases/{kb_id}")
def delete_kb(kb_id: int, db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    db.delete(kb)
    db.commit()
    return {"code": 200, "data": {"id": kb_id}}

@app.post("/api/ai/knowledge-bases/{kb_id}/activate")
def activate_kb(kb_id: int, db: Session = Depends(get_db)):
    db.query(KnowledgeBaseConfig).update({KnowledgeBaseConfig.is_active: 0})
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    kb.is_active = 1
    db.commit()
    return {"code": 200, "data": {"id": kb_id}}

# --- 知识库文档管理（RAG） ---

@app.post("/api/ai/knowledge-bases/{kb_id}/documents")
async def upload_document(
    kb_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """上传知识库文档（PDF/DOCX/TXT），自动处理为RAG可用的分块"""
    kb = db.query(KnowledgeBaseConfig).filter(KnowledgeBaseConfig.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")

    # 验证文件类型
    filename = file.filename or "unknown"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("pdf", "docx", "doc", "txt"):
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: .{ext}，仅支持 PDF/DOCX/TXT")

    # 保存文件
    file_content = await file.read()
    stored_filename = f"{uuid.uuid4().hex}.{ext}"
    file_path = os.path.join(KB_DOCS_DIR, stored_filename)
    with open(file_path, "wb") as f:
        f.write(file_content)

    # 创建文档记录
    doc = KnowledgeDocument(
        kb_id=kb_id,
        filename=filename,
        file_path=file_path,
        file_type=ext,
        file_size=len(file_content),
        status="processing",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    # 获取模型API配置（用于生成嵌入向量）
    active_model = db.query(AIModelConfig).filter(AIModelConfig.is_active == 1).first()
    emb_api_key = active_model.api_key if active_model else ""
    emb_base_url = active_model.base_url if active_model else ""

    # 异步处理文档（提取文本→分块→嵌入→存储）
    try:
        await process_document(
            file_path=file_path,
            file_type=ext,
            doc_id=doc.id,
            kb_id=kb_id,
            db_session=db,
            api_key=emb_api_key,
            base_url=emb_base_url,
        )
    except Exception as e:
        # 处理失败不影㘿上传，记录错误
        print(f"[文档上传] 处理失败: {e}")

    # 重新查询获取最新状态
    db.refresh(doc)
    return {"code": 200, "data": {
        "id": doc.id, "filename": doc.filename, "file_type": doc.file_type,
        "file_size": doc.file_size, "chunk_count": doc.chunk_count,
        "status": doc.status, "error_message": doc.error_message,
    }}


@app.get("/api/ai/knowledge-bases/{kb_id}/documents")
def list_documents(kb_id: int, db: Session = Depends(get_db)):
    """获取知识库下的所有文档"""
    docs = db.query(KnowledgeDocument).filter(
        KnowledgeDocument.kb_id == kb_id
    ).order_by(KnowledgeDocument.created_at.desc()).all()
    
    return {"code": 200, "data": [{
        "id": d.id, "filename": d.filename, "file_type": d.file_type,
        "file_size": d.file_size, "chunk_count": d.chunk_count,
        "status": d.status, "error_message": d.error_message,
        "created_at": str(d.created_at),
    } for d in docs]}


@app.delete("/api/ai/knowledge-bases/{kb_id}/documents/{doc_id}")
def delete_document(kb_id: int, doc_id: int, db: Session = Depends(get_db)):
    """删除知识库文档及其分块"""
    doc = db.query(KnowledgeDocument).filter(
        KnowledgeDocument.id == doc_id,
        KnowledgeDocument.kb_id == kb_id,
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")

    # 删除分块记录
    db.query(DocumentChunk).filter(DocumentChunk.doc_id == doc_id).delete()

    # 删除物理文件
    if os.path.exists(doc.file_path):
        os.remove(doc.file_path)

    # 删除文档记录
    db.delete(doc)
    db.commit()
    return {"code": 200, "data": {"id": doc_id}}

# --- 分析技能 ---

@app.get("/api/ai/skills")
def list_skills(db: Session = Depends(get_db)):
    skills = db.query(AnalysisSkill).order_by(AnalysisSkill.created_at.desc()).all()
    return {"code": 200, "data": [{"id": s.id, "name": s.name, "description": s.description,
        "is_active": bool(s.is_active), "created_at": str(s.created_at)} for s in skills]}

@app.post("/api/ai/skills")
def create_skill(req: SkillCreate, db: Session = Depends(get_db)):
    skill = AnalysisSkill(**req.dict())
    db.add(skill)
    db.commit()
    db.refresh(skill)
    return {"code": 200, "data": {"id": skill.id, "name": skill.name}}

@app.get("/api/ai/skills/{skill_id}")
def get_skill(skill_id: int, db: Session = Depends(get_db)):
    skill = db.query(AnalysisSkill).filter(AnalysisSkill.id == skill_id).first()
    if not skill:
        raise HTTPException(status_code=404, detail="分析技能不存在")
    return {"code": 200, "data": {"id": skill.id, "name": skill.name, "description": skill.description,
        "system_prompt": skill.system_prompt, "user_prompt_template": skill.user_prompt_template,
        "is_active": bool(skill.is_active)}}

@app.put("/api/ai/skills/{skill_id}")
def update_skill(skill_id: int, req: SkillCreate, db: Session = Depends(get_db)):
    skill = db.query(AnalysisSkill).filter(AnalysisSkill.id == skill_id).first()
    if not skill:
        raise HTTPException(status_code=404, detail="分析技能不存在")
    for k, v in req.dict().items():
        setattr(skill, k, v)
    db.commit()
    return {"code": 200, "data": {"id": skill.id}}

@app.delete("/api/ai/skills/{skill_id}")
def delete_skill(skill_id: int, db: Session = Depends(get_db)):
    skill = db.query(AnalysisSkill).filter(AnalysisSkill.id == skill_id).first()
    if not skill:
        raise HTTPException(status_code=404, detail="分析技能不存在")
    db.delete(skill)
    db.commit()
    return {"code": 200, "data": {"id": skill_id}}

@app.post("/api/ai/skills/{skill_id}/activate")
def activate_skill(skill_id: int, db: Session = Depends(get_db)):
    db.query(AnalysisSkill).update({AnalysisSkill.is_active: 0})
    skill = db.query(AnalysisSkill).filter(AnalysisSkill.id == skill_id).first()
    if not skill:
        raise HTTPException(status_code=404, detail="分析技能不存在")
    skill.is_active = 1
    db.commit()
    return {"code": 200, "data": {"id": skill_id}}


# ========== 健康检查 ==========

@app.get("/api/health")
def health_check():
    return {"code": 200, "status": "ok", "service": "监理巡视闭环智能体"}


# ========== 前端静态文件托管 ==========
# 必须放在所有API路由之后，确保 /api/* 优先匹配
STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")

@app.get("/{full_path:path}")
async def serve_frontend(full_path: str):
    """托管前端静态文件，支持 .html 回退和 SPA 路由"""
    if not os.path.isdir(STATIC_DIR):
        raise HTTPException(status_code=404, detail="Not Found")

    # 1. 精确文件匹配（JS/CSS/图片等静态资源）
    file_path = os.path.join(STATIC_DIR, full_path)
    if os.path.isfile(file_path):
        return FileResponse(file_path)

    # 2. 目录 → index.html
    if os.path.isdir(file_path):
        index_path = os.path.join(file_path, "index.html")
        if os.path.isfile(index_path):
            return FileResponse(index_path)

    # 3. 路径.html 回退（/input → input.html）
    html_path = os.path.join(STATIC_DIR, f"{full_path}.html")
    if os.path.isfile(html_path):
        return FileResponse(html_path)

    # 4. SPA 回退 → index.html
    fallback = os.path.join(STATIC_DIR, "index.html")
    if os.path.isfile(fallback):
        return FileResponse(fallback)

    raise HTTPException(status_code=404, detail="Not Found")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
